"""
build_benchmark_sheet.py: builds benchmark_sheet.xlsx from benchmark_cases.json.
 
Re-run if the cases change:  python build_benchmark_sheet.py
Then recalculate formulas (LibreOffice) so cached values exist.
"""
 
import json
from pathlib import Path
 
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
 
HERE = Path(__file__).parent
cases = json.loads((HERE / "benchmark_cases.json").read_text(encoding="utf-8"))["cases"]
 
FONT = "Arial"
BLUE = Font(name=FONT, size=10, color="0000FF")          # inputs
BLACK = Font(name=FONT, size=10, color="000000")         # formulas
GREEN = Font(name=FONT, size=10, color="008000")         # cross-sheet links
BOLD = Font(name=FONT, size=10, bold=True)
HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True)
YELLOW = PatternFill("solid", start_color="FFFF00")
NAVY = PatternFill("solid", start_color="1F3864")
GREY = PatternFill("solid", start_color="EDEDED")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
 
 
def style_header(ws, row, first_col, last_col):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD, NAVY, CENTER, BOX
 
 
def money(s):
    return f"${float(s):,.2f}"
 
 
def expected_text(c):
    e = c["expected"]
    s = c["scenario"]
    if s == "auto_loan":
        return (f"Payment {money(e['monthly_payment'])}/mo; total interest {money(e['total_interest'])}; "
                f"total cost {money(e['total_cost'])}; cash at signing {money(e['cash_due_at_signing'])}; "
                f"true APR {e['true_apr_percent']}%")
    if s == "auto_loan_compare":
        return (f"Option A total {money(e['a_total_cost'])}; Option B total {money(e['b_total_cost'])}; "
                f"cheaper: {e['cheaper']} by {money(e['difference'])}")
    if s == "lease":
        return (f"Monthly (pre-tax) {money(e['monthly_payment_pre_tax'])}; total lease cost "
                f"{money(e['total_lease_cost'])}; approx APR {e['approx_apr_percent']}%; "
                f"effective cost/month {money(e['effective_cost_per_month'])}")
    if s == "bnpl":
        return (f"First payment {money(e['first_payment'])}; later payments {money(e['later_payment_each'])} each; "
                f"fees {money(e['fees_total'])}; total {money(e['total_cost'])}; "
                f"extra vs cash {money(e['extra_cost_vs_cash'])}")
    if s == "card":
        return (f"{e['months_to_payoff']} months to pay off; total interest {money(e['total_interest'])}; "
                f"total paid {money(e['total_paid'])}")
    return ""
 
 
wb = Workbook()
 
# ---------------------------------------------------------------- HowToUse
ws = wb.active
ws.title = "HowToUse"
ws["A1"] = "Benchmark Sheet: Our Tool vs. a General Chatbot"
ws["A1"].font = TITLE
rows = [
    ("Purpose", "Score how well our tool and a general chatbot handle the same 12 realistic offers. This sheet becomes the proof section of the pitch."),
    ("Step 1", "On the Cases tab, copy each 'Offer text' exactly as written into a general chatbot (same wording every time) and into our tool. Save both outputs."),
    ("Step 2", "On the Scoring tab, fill the yellow cells with Yes or No (dropdown). Do not edit anything else."),
    ("Numbers correct", "Yes only if every key number matches the 'Correct answers' on the Cases tab (within one cent for exact figures; chatbot answers that round to the dollar count as correct if the rounding is right)."),
    ("Hidden item caught", "Yes if the output clearly points out the 'Hidden item' on the Cases tab. For CONTROL cases (A1, B1) there is nothing hidden: mark Yes if the tool correctly raised NO alarm, and No if it invented a problem."),
    ("Sources cited", "Yes if factual claims (not just math) come with a named, checkable source."),
    ("Honesty rule", "Record results exactly as they happen, including where our tool loses. Judges trust honest numbers. Paste unedited outputs into a separate doc and link it in notes."),
    ("Where answers come from", "Correct answers were computed independently of our calculators (see make_benchmark_cases.py) and live in benchmark_cases.json. Run 'python run_benchmark.py' to check the math layer automatically."),
    ("Colour key", "Blue text = input you type. Black = formula. Green = link to another tab. Yellow fill = cells you fill in."),
]
r = 3
for k, v in rows:
    ws.cell(row=r, column=1, value=k).font = BOLD
    c = ws.cell(row=r, column=2, value=v)
    c.font, c.alignment = BLACK, WRAP
    ws.cell(row=r, column=1).alignment = WRAP
    r += 1
r += 1
ws.cell(row=r, column=1, value="Example row (format only; not counted)").font = BOLD
r += 1
for col, h in enumerate(["Case", "Numbers correct", "Hidden item caught", "Sources cited"], start=1):
    ws.cell(row=r, column=col, value=h)
style_header(ws, r, 1, 4)
r += 1
for col, v in enumerate(["A2", "Yes", "Yes", "No"], start=1):
    c = ws.cell(row=r, column=col, value=v)
    c.font, c.border = BLUE, BOX
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 110
ws.column_dimensions["C"].width = 20
ws.column_dimensions["D"].width = 16
 
# ---------------------------------------------------------------- Cases
wc = wb.create_sheet("Cases")
wc["A1"] = "Benchmark cases (reference data; source: benchmark_cases.json)"
wc["A1"].font = TITLE
wc["A2"] = ("Correct answers were computed independently of calculators.py. "
            "Control = a fair offer where the tool should NOT raise alarms.")
wc["A2"].font = Font(name=FONT, size=9, italic=True)
hdr = ["Case", "Scenario", "Control?", "Title", "Offer text (paste this)", "Correct answers", "Hidden item to catch"]
HR = 4
for i, h in enumerate(hdr, start=1):
    wc.cell(row=HR, column=i, value=h)
style_header(wc, HR, 1, len(hdr))
for i, c in enumerate(cases):
    row = HR + 1 + i
    vals = [c["id"], c["scenario"], "Yes" if c["control"] else "No", c["title"],
            c["offer_text"], expected_text(c), c["hidden_item"]]
    for col, v in enumerate(vals, start=1):
        cell = wc.cell(row=row, column=col, value=v)
        cell.font, cell.alignment, cell.border = BLUE if col in (5, 6, 7) else BLACK, WRAP, BOX
    wc.row_dimensions[row].height = 78
for col, w in zip("ABCDEFG", (7, 17, 10, 30, 52, 52, 48)):
    wc.column_dimensions[col].width = w
wc.freeze_panes = "B5"
 
# ---------------------------------------------------------------- Scoring
sc = wb.create_sheet("Scoring")
sc["A1"] = "Scoring: fill the yellow cells with Yes or No"
sc["A1"].font = TITLE
sc.merge_cells("D3:F3")
sc.merge_cells("G3:I3")
sc["D3"], sc["G3"] = "OUR TOOL", "GENERAL CHATBOT"
for ref in ("D3", "G3"):
    sc[ref].font, sc[ref].fill, sc[ref].alignment = HEAD, NAVY, CENTER
sh = ["Case", "Scenario", "Control?",
      "Numbers correct", "Hidden item caught", "Sources cited",
      "Numbers correct", "Hidden item caught", "Sources cited", "Notes"]
SR = 4
for i, h in enumerate(sh, start=1):
    sc.cell(row=SR, column=i, value=h)
style_header(sc, SR, 1, len(sh))
FIRST, LAST = SR + 1, SR + len(cases)
for i, c in enumerate(cases):
    row = FIRST + i
    src = HR + 1 + i
    for col, srccol in ((1, "A"), (2, "B"), (3, "C")):
        cell = sc.cell(row=row, column=col, value=f"=Cases!{srccol}{src}")
        cell.font, cell.border, cell.alignment = GREEN, BOX, CENTER
    for col in range(4, 11):
        cell = sc.cell(row=row, column=col)
        cell.fill, cell.border = YELLOW, BOX
        cell.font = BLUE
        cell.alignment = CENTER if col < 10 else WRAP
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
sc.add_data_validation(dv)
dv.add(f"D{FIRST}:I{LAST}")
for col, w in zip("ABCDEFGHIJ", (8, 18, 10, 14, 14, 14, 14, 14, 14, 44)):
    sc.column_dimensions[col].width = w
sc.freeze_panes = f"D{FIRST}"
 
# ---------------------------------------------------------------- Summary
sm = wb.create_sheet("Summary")
sm["A1"] = "Summary (all values are formulas; they update as you score)"
sm["A1"].font = TITLE
for i, h in enumerate(["Measure", "Our tool", "General chatbot", "Difference (ours - chatbot)"], start=1):
    sm.cell(row=3, column=i, value=h)
style_header(sm, 3, 1, 4)
 
sm["A4"] = "Total cases"
sm["B4"] = f"=COUNTA(Scoring!A{FIRST}:A{LAST})"
sm["C4"] = "=B4"
 
measures = [("Numbers correct", "D", "G"), ("Hidden item caught", "E", "H"), ("Sources cited", "F", "I")]
row = 5
count_rows = {}
for name, ours, theirs in measures:
    sm.cell(row=row, column=1, value=f"{name} (count of Yes)")
    sm.cell(row=row, column=2, value=f'=COUNTIF(Scoring!{ours}{FIRST}:{ours}{LAST},"Yes")')
    sm.cell(row=row, column=3, value=f'=COUNTIF(Scoring!{theirs}{FIRST}:{theirs}{LAST},"Yes")')
    sm.cell(row=row, column=4, value=f"=B{row}-C{row}")
    count_rows[name] = row
    row += 1
for name, ours, theirs in measures:
    sm.cell(row=row, column=1, value=f"{name} (% of cases)")
    src_row = count_rows[name]
    sm.cell(row=row, column=2, value=f"=IF($B$4=0,0,B{src_row}/$B$4)")
    sm.cell(row=row, column=3, value=f"=IF($C$4=0,0,C{src_row}/$C$4)")
    sm.cell(row=row, column=4, value=f"=B{row}-C{row}")
    for col in (2, 3, 4):
        sm.cell(row=row, column=col).number_format = "0.0%"
    row += 1
sm.cell(row=row, column=1, value="Rows scored so far (our tool / chatbot, out of 3 columns each)")
sm.cell(row=row, column=2, value=f"=COUNTA(Scoring!D{FIRST}:F{LAST})")
sm.cell(row=row, column=3, value=f"=COUNTA(Scoring!G{FIRST}:I{LAST})")
scored_row = row
row += 1
sm.cell(row=row, column=1, value="Fully scored?")
sm.cell(row=row, column=2, value=f'=IF(B{scored_row}=$B$4*3,"Yes","Not yet")')
sm.cell(row=row, column=3, value=f'=IF(C{scored_row}=$C$4*3,"Yes","Not yet")')
row += 2
 
sm.cell(row=row, column=1, value="Lines for Devpost / pitch (build automatically)").font = BOLD
row += 1
nr, hr_, sr_ = count_rows["Numbers correct"], count_rows["Hidden item caught"], count_rows["Sources cited"]
lines = [
    f'="Numeric answers correct: ours "&B{nr}&" of "&B4&" vs. chatbot "&C{nr}&" of "&C4',
    f'="Hidden costs caught (and false alarms avoided): ours "&B{hr_}&" of "&B4&" vs. chatbot "&C{hr_}&" of "&C4',
    f'="Claims with named sources: ours "&B{sr_}&" of "&B4&" vs. chatbot "&C{sr_}&" of "&C4',
]
for ln in lines:
    c = sm.cell(row=row, column=1, value=ln)
    row += 1
 
for r_ in range(4, row):
    for col in range(1, 5):
        cell = sm.cell(row=r_, column=col)
        if cell.value is not None and cell.font != BOLD and cell.font != HEAD:
            cell.font = BLACK
        if col > 1 and r_ <= scored_row + 1:
            cell.alignment = Alignment(horizontal="center")
        if r_ <= scored_row + 1:
            cell.border = BOX
sm["A4"].font = BOLD
sm.column_dimensions["A"].width = 62
for col in "BCD":
    sm.column_dimensions[col].width = 26
 
out = HERE / "benchmark_sheet.xlsx"
wb.save(out)
print("saved", out)